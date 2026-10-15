"""从「联系单看板」Excel 导入/更新联系单与项目数据。

支持 SQLite 和 MySQL（通过 SQLAlchemy Engine）。
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import Engine, create_engine, text
from sqlalchemy.engine import RowMapping

from steeltech_db.project_split import is_processing_contact_id, resolve_import_project_no

CONTACT_STATUS_MAP = {
    "未开始": "pending",
    "进行中": "processing",
    "已完成": "done",
    "已取消": "cancelled",
    "项目暂停": "processing",
}

PROJECT_STATUS_MAP = {
    "未开始": "not_started",
    "进行中": "active",
    "已完成": "done",
    "已取消": "cancelled",
    "项目暂停": "paused",
}

CONTACT_STATUS_PRIORITY = {
    "进行中": 5,
    "未开始": 4,
    "项目暂停": 3,
    "已完成": 2,
    "已取消": 1,
}

DESIGN_CATEGORIES = frozenset({"设计", "设计问题处理", "安全评估"})
DETAIL_CATEGORIES = frozenset({"钢结构细化", "加工单"})
DETAIL_ISSUE_CATEGORIES = frozenset({"细化问题处理"})
# 楼层板排版 = 排楼承板（排楼板），与 plate_layout（钢板排板）不同
FLOOR_DECK_LAYOUT_CATEGORIES = frozenset({"楼层板排版"})
TILE_LAYOUT_CATEGORIES = frozenset({"瓦片排版"})

DTP_PATTERN = re.compile(r"DTP\d+", re.IGNORECASE)
PROJECT_PREFIX_PATTERN = re.compile(r"^【[^】]*】")
PROJECT_NO_PATTERN = re.compile(r"^[A-Z]{0,3}\d+", re.IGNORECASE)


def normalize_relative_path(value: str) -> str:
    trimmed = (value or "").strip()
    if not trimmed:
        return ""

    unc_match = re.match(r"^\\\\[^\\]+\\([A-Za-z])\$?\\(.*)$", trimmed, re.IGNORECASE)
    if unc_match and unc_match.group(2) is not None:
        share = unc_match.group(1).lower()
        rest = unc_match.group(2).replace("/", "\\").lstrip("\\")
        return f"{share}\\{rest}" if rest else share

    ip_path_match = re.match(r"^[\d.]+\\([A-Za-z])\\(.*)$", trimmed, re.IGNORECASE)
    if ip_path_match and ip_path_match.group(2) is not None:
        share = ip_path_match.group(1).lower()
        rest = ip_path_match.group(2).replace("/", "\\").lstrip("\\")
        return f"{share}\\{rest}" if rest else share

    drive_match = re.match(r"^([A-Za-z]):[\\/]*(.*)$", trimmed, re.IGNORECASE)
    if drive_match and drive_match.group(1):
        share = drive_match.group(1).lower()
        rest = (drive_match.group(2) or "").replace("/", "\\").lstrip("\\")
        return f"{share}\\{rest}" if rest else share

    return trimmed.replace("/", "\\").lstrip("\\")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def normalize_project_no(raw: Any) -> str:
    text = _cell_text(raw)
    if not text:
        return ""
    text = PROJECT_PREFIX_PATTERN.sub("", text).strip()
    match = PROJECT_NO_PATTERN.match(text)
    return match.group(0).upper() if match else text.upper()


def parse_natures(category: str) -> list[str]:
    if not category:
        return []

    natures: list[str] = []
    parts = [part.strip() for part in category.split(",") if part.strip()]
    for part in parts:
        if part in DESIGN_CATEGORIES and "design" not in natures:
            natures.append("design")
        if part in DETAIL_CATEGORIES and "detail" not in natures:
            natures.append("detail")
        if part in DETAIL_ISSUE_CATEGORIES and "detail_issue" not in natures:
            natures.append("detail_issue")
        if part in TILE_LAYOUT_CATEGORIES and "tile_layout" not in natures:
            natures.append("tile_layout")
        if part in FLOOR_DECK_LAYOUT_CATEGORIES and "floor_deck_layout" not in natures:
            natures.append("floor_deck_layout")
    return natures


def parse_person_names(raw: Any) -> list[str]:
    text = _cell_text(raw)
    if not text:
        return []
    return [name.strip() for name in re.split(r"[,，、]", text) if name.strip()]


def build_content(row: dict[str, Any]) -> str:
    lines: list[str] = []
    mapping = [
        ("category", "项目分类"),
        ("structure_type", "结构形式"),
        ("workload", "工程量"),
        ("remark", "备注"),
        ("path", "路径"),
        ("owners", "项目负责人"),
    ]
    for key, label in mapping:
        value = _cell_text(row.get(key))
        if value:
            lines.append(f"{label}：{value}")
    return "\n".join(lines)


JIAGONGDAN_CONTACT_LABEL = "加工单"


def extract_contact_id(raw: Any) -> str | None:
    text = _cell_text(raw)
    if not text or text == "/":
        return None
    if text == JIAGONGDAN_CONTACT_LABEL:
        return JIAGONGDAN_CONTACT_LABEL

    dtps = DTP_PATTERN.findall(text.upper())
    if dtps:
        return dtps[0].upper()

    if re.fullmatch(r"[A-Z0-9]+", text.upper()):
        return text.upper()
    return None


def _execute(engine: Engine, sql: str, params: dict | None = None):
    """执行 SQL 并返回结果代理。"""
    with engine.connect() as conn:
        result = conn.execute(text(sql), params or {})
        return result


def _execute_scalar(engine: Engine, sql: str, params: dict | None = None):
    with engine.connect() as conn:
        return conn.execute(text(sql), params or {}).fetchone()


def _execute_all(engine: Engine, sql: str, params: dict | None = None) -> list[RowMapping]:
    with engine.connect() as conn:
        return list(conn.execute(text(sql), params or {}).mappings().all())


def _execute_commit(engine: Engine, sql: str, params: dict | None = None) -> None:
    with engine.connect() as conn:
        trans = conn.begin()
        try:
            conn.execute(text(sql), params or {})
            trans.commit()
        except Exception:
            trans.rollback()
            raise


def _insert_ignore_sql(engine: Engine, table: str, columns: str) -> str:
    """跨数据库 INSERT IGNORE 语句前缀。"""
    if engine.dialect.name == "mysql":
        return f"INSERT IGNORE INTO {table} ({columns})"
    return f"INSERT OR IGNORE INTO {table} ({columns})"


def find_existing_processing_contact(engine: Engine, project_no: str) -> str | None:
    rows = _execute_all(
        engine,
        "SELECT contact_form_id FROM contact_form_projects WHERE project_no = :project_no",
        {"project_no": project_no},
    )
    for row in rows:
        normalized = str(row["contact_form_id"] or "").strip()
        if is_processing_contact_id(normalized):
            return normalized
    return None


def allocate_brd_id(engine: Engine, project_no: str, received_date: str) -> str:
    row = _execute_scalar(
        engine,
        "SELECT contact_form_id FROM contact_form_projects WHERE project_no = :project_no LIMIT 1",
        {"project_no": project_no},
    )
    if row:
        return str(row[0])

    digits = re.sub(r"^[A-Z]+", "", project_no)
    yymmdd = received_date[2:4] + received_date[5:7] + received_date[8:10]
    prefix = f"BRD{yymmdd}C{digits}"
    rows = _execute_all(
        engine,
        "SELECT id FROM contact_forms WHERE id LIKE :pattern",
        {"pattern": prefix + "%"},
    )
    max_seq = 0
    for row in rows:
        contact_id = str(row["id"])
        match = re.fullmatch(rf"{re.escape(prefix)}(\d{{4}})", contact_id)
        if match:
            max_seq = max(max_seq, int(match.group(1)))
    return f"{prefix}{max_seq + 1:04d}"


def resolve_contact_id(
    engine: Engine,
    raw_contact_id: Any,
    project_no: str,
    received_date: str,
) -> str:
    contact_id = extract_contact_id(raw_contact_id)
    if contact_id == JIAGONGDAN_CONTACT_LABEL:
        existing = find_existing_processing_contact(engine, project_no)
        if existing:
            return existing
        return f"{JIAGONGDAN_CONTACT_LABEL}-{project_no}"
    if contact_id:
        return contact_id
    return allocate_brd_id(engine, project_no, received_date)


@dataclass
class BoardRow:
    received_date: str
    category: str
    project_no: str
    project_name: str
    raw_contact_id: str
    structure_type: str
    workload: str
    excel_status: str
    remark: str
    path: str
    owners: str
    contact_id: str = ""
    import_project_no: str = ""

    def to_content_dict(self) -> dict[str, Any]:
        return {
            "category": self.category,
            "structure_type": self.structure_type,
            "workload": self.workload,
            "remark": self.remark,
            "path": self.path,
            "owners": self.owners,
        }


@dataclass
class ImportStats:
    rows_read: int = 0
    rows_skipped: int = 0
    projects_upserted: int = 0
    contacts_upserted: int = 0
    links_synced: int = 0
    nature_links: int = 0
    personnel_links: int = 0
    warnings: list[str] = field(default_factory=list)


def read_contact_board_rows(excel_path: Path) -> list[BoardRow]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    headers = [(_cell_text(cell) or f"col_{index}") for index, cell in enumerate(next(rows))]

    column_index = {name: idx for idx, name in enumerate(headers)}
    required = ["收单时间", "项目分类", "项目编号", "项目名称", "联系单编号", "状态"]
    missing = [name for name in required if name not in column_index]
    if missing:
        raise ValueError(f"Excel 缺少必要列: {', '.join(missing)}")

    def get_value(record: tuple[Any, ...], name: str) -> Any:
        idx = column_index.get(name)
        if idx is None:
            return ""
        return record[idx] if idx < len(record) else ""

    parsed: list[BoardRow] = []
    for record in rows:
        project_no = normalize_project_no(get_value(record, "项目编号"))
        project_name = _cell_text(get_value(record, "项目名称"))
        if not project_no or not project_name:
            continue
        parsed.append(
            BoardRow(
                received_date=_cell_text(get_value(record, "收单时间")),
                category=_cell_text(get_value(record, "项目分类")),
                project_no=project_no,
                project_name=project_name,
                raw_contact_id=_cell_text(get_value(record, "联系单编号")),
                structure_type=_cell_text(get_value(record, "结构形式")),
                workload=_cell_text(get_value(record, "工程量")),
                excel_status=_cell_text(get_value(record, "状态")),
                remark=_cell_text(get_value(record, "备注")),
                path=_cell_text(get_value(record, "路径")),
                owners=_cell_text(get_value(record, "项目负责人")),
            )
        )
    workbook.close()
    return parsed


def _pick_primary_row(rows: list[BoardRow]) -> BoardRow:
    def sort_key(row: BoardRow) -> tuple[int, str]:
        return (
            CONTACT_STATUS_PRIORITY.get(row.excel_status, 0),
            row.received_date,
        )

    return sorted(rows, key=sort_key, reverse=True)[0]


def _aggregate_contact_status(rows: list[BoardRow]) -> str:
    best = max(
        rows,
        key=lambda row: CONTACT_STATUS_PRIORITY.get(row.excel_status, 0),
    )
    return CONTACT_STATUS_MAP.get(best.excel_status, "pending")


def import_contact_board(
    excel_path: Path,
    db_path: Path,
    *,
    now: str | None = None,
) -> ImportStats:
    """从联系单看板 Excel 导入数据（SQLite 模式，向后兼容）。

    如需 MySQL，请使用 import_contact_board_to_engine() 并传入 MySQL Engine。
    """
    engine = create_engine(f"sqlite:///{db_path.as_posix()}")
    try:
        return _import_contact_board_impl(excel_path, engine, now=now)
    finally:
        engine.dispose()


def import_contact_board_to_engine(
    excel_path: Path,
    engine: Engine,
    *,
    now: str | None = None,
) -> ImportStats:
    """从联系单看板 Excel 导入数据到指定数据库 Engine（跨数据库兼容）。"""
    return _import_contact_board_impl(excel_path, engine, now=now)


def _import_contact_board_impl(
    excel_path: Path,
    engine: Engine,
    *,
    now: str | None = None,
) -> ImportStats:
    """核心导入逻辑 —— 跨数据库（SQLite / MySQL）兼容。"""
    stats = ImportStats()
    now_text = now or datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    board_rows = read_contact_board_rows(excel_path)
    stats.rows_read = len(board_rows)
    if not board_rows:
        return stats

    personnel_by_name = {
        str(row["name"]): str(row["id"])
        for row in _execute_all(engine, "SELECT id, name FROM personnel")
    }

    for row in board_rows:
        if not row.received_date:
            row.received_date = now_text[:10]
        row.contact_id = resolve_contact_id(
            engine,
            row.raw_contact_id,
            row.project_no,
            row.received_date,
        )

    projects_by_no: dict[str, BoardRow] = {}
    base_primary_contact: dict[str, str] = {}
    for row in board_rows:
        row.import_project_no = resolve_import_project_no(
            row.project_no,
            row.contact_id,
            base_primary_contact=base_primary_contact,
        )
        existing = projects_by_no.get(row.import_project_no)
        if existing is None or row.received_date >= existing.received_date:
            projects_by_no[row.import_project_no] = row

    contacts_by_id: dict[str, list[BoardRow]] = defaultdict(list)
    links: set[tuple[str, str]] = set()
    for row in board_rows:
        contacts_by_id[row.contact_id].append(row)
        links.add((row.contact_id, row.import_project_no))

    with engine.connect() as conn:
        trans = conn.begin()
        try:
            for project_no, row in projects_by_no.items():
                local_path = normalize_relative_path(row.path)
                project_status = PROJECT_STATUS_MAP.get(row.excel_status, "active")
                natures = parse_natures(row.category)
                design_completed = int("design" in natures and row.excel_status == "已完成")
                detail_completed = int("detail" in natures and row.excel_status == "已完成")
                exists = conn.execute(
                    text("SELECT 1 FROM projects WHERE project_no = :project_no"),
                    {"project_no": project_no},
                ).fetchone()
                if exists:
                    conn.execute(
                        text(
                            """
                            UPDATE projects
                            SET name = :name, status = :status,
                                design_completed = :design_completed,
                                detail_completed = :detail_completed,
                                received_date = :received_date,
                                local_work_path = :local_work_path,
                                updated_at = :updated_at
                            WHERE project_no = :project_no
                            """
                        ),
                        {
                            "name": row.project_name,
                            "status": project_status,
                            "design_completed": design_completed,
                            "detail_completed": detail_completed,
                            "received_date": row.received_date,
                            "local_work_path": local_path,
                            "updated_at": now_text,
                            "project_no": project_no,
                        },
                    )
                else:
                    conn.execute(
                        text(
                            """
                            INSERT INTO projects (
                              project_no, name, customer, status,
                              design_completed, detail_completed,
                              received_date, local_work_path,
                              created_at, updated_at
                            ) VALUES (
                              :project_no, :name, NULL, :status,
                              :design_completed, :detail_completed,
                              :received_date, :local_work_path,
                              :created_at, :updated_at
                            )
                            """
                        ),
                        {
                            "project_no": project_no,
                            "name": row.project_name,
                            "status": project_status,
                            "design_completed": design_completed,
                            "detail_completed": detail_completed,
                            "received_date": row.received_date,
                            "local_work_path": local_path,
                            "created_at": now_text,
                            "updated_at": now_text,
                        },
                    )
                stats.projects_upserted += 1

                conn.execute(
                    text("DELETE FROM project_natures WHERE project_no = :project_no"),
                    {"project_no": project_no},
                )
                insert_nature_sql = _insert_ignore_sql(engine, "project_natures", "project_no, nature, created_at")
                for nature in natures:
                    conn.execute(
                        text(f"{insert_nature_sql} VALUES (:project_no, :nature, :created_at)"),
                        {"project_no": project_no, "nature": nature, "created_at": now_text},
                    )
                    stats.nature_links += 1

                conn.execute(
                    text("DELETE FROM project_personnel WHERE project_no = :project_no"),
                    {"project_no": project_no},
                )
                insert_personnel_sql = _insert_ignore_sql(engine, "project_personnel", "project_no, personnel_id, created_at")
                for owner_name in parse_person_names(row.owners):
                    personnel_id = personnel_by_name.get(owner_name)
                    if personnel_id is None:
                        stats.warnings.append(
                            f"未找到项目负责人「{owner_name}」（项目 {project_no} / {row.project_no}）"
                        )
                        continue
                    conn.execute(
                        text(f"{insert_personnel_sql} VALUES (:project_no, :personnel_id, :created_at)"),
                        {"project_no": project_no, "personnel_id": personnel_id, "created_at": now_text},
                    )
                    stats.personnel_links += 1

            excel_project_nos = set(projects_by_no)
            for project_no in excel_project_nos:
                conn.execute(
                    text("DELETE FROM contact_form_projects WHERE project_no = :project_no"),
                    {"project_no": project_no},
                )

            for contact_id, grouped_rows in contacts_by_id.items():
                primary = _pick_primary_row(grouped_rows)
                contact_status = _aggregate_contact_status(grouped_rows)
                content = build_content(primary.to_content_dict())
                created_at = f"{primary.received_date} 09:00:00"
                exists = conn.execute(
                    text("SELECT 1 FROM contact_forms WHERE id = :id"),
                    {"id": contact_id},
                ).fetchone()
                if exists:
                    conn.execute(
                        text(
                            """
                            UPDATE contact_forms
                            SET title = :title, received_date = :received_date,
                                urgency = :urgency, status = :status,
                                content = :content, root_id = :root_id,
                                relation_type = :relation_type,
                                sort_order = :sort_order,
                                updated_at = :updated_at
                            WHERE id = :id
                            """
                        ),
                        {
                            "title": primary.project_name,
                            "received_date": primary.received_date,
                            "urgency": "普通",
                            "status": contact_status,
                            "content": content,
                            "root_id": contact_id,
                            "relation_type": "primary",
                            "sort_order": 0,
                            "updated_at": now_text,
                            "id": contact_id,
                        },
                    )
                else:
                    conn.execute(
                        text(
                            """
                            INSERT INTO contact_forms (
                              id, title, received_date, urgency, status, content,
                              expect_reply_date, parent_id, root_id, relation_type,
                              sort_order, created_at, updated_at
                            ) VALUES (
                              :id, :title, :received_date, :urgency, :status, :content,
                              '', NULL, :root_id, 'primary', 0, :created_at, :updated_at
                            )
                            """
                        ),
                        {
                            "id": contact_id,
                            "title": primary.project_name,
                            "received_date": primary.received_date,
                            "urgency": "普通",
                            "status": contact_status,
                            "content": content,
                            "root_id": contact_id,
                            "created_at": created_at,
                            "updated_at": now_text,
                        },
                    )
                stats.contacts_upserted += 1

                insert_link_sql = _insert_ignore_sql(
                    engine,
                    "contact_form_projects",
                    "contact_form_id, project_no, source_type, source_contact_form_id, created_at",
                )
                for link_contact_id, project_no in links:
                    if link_contact_id != contact_id:
                        continue
                    conn.execute(
                        text(
                            f"{insert_link_sql} VALUES "
                            "(:contact_form_id, :project_no, 'own', NULL, :created_at)"
                        ),
                        {
                            "contact_form_id": contact_id,
                            "project_no": project_no,
                            "created_at": now_text,
                        },
                    )
                    stats.links_synced += 1

            conn.execute(
                text(
                    """
                    DELETE FROM contact_forms
                    WHERE id LIKE 'BRD%'
                      AND id NOT IN (SELECT contact_form_id FROM contact_form_projects)
                    """
                )
            )
            trans.commit()
        except Exception:
            trans.rollback()
            raise

    return stats
