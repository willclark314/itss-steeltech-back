"""钢结构技术科结算业务逻辑"""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

from steeltech_db.extensions import db
from steeltech_db.models import SettlementPendingItem, SettlementPricingBasis, SettlementRecord

_CATEGORY_DETAIL = "detail"
_CATEGORY_DESIGN = "design"
_CATEGORY_LABELS = {
    "细化": _CATEGORY_DETAIL,
    "设计": _CATEGORY_DESIGN,
    _CATEGORY_DETAIL: _CATEGORY_DETAIL,
    _CATEGORY_DESIGN: _CATEGORY_DESIGN,
}
_CATEGORY_DISPLAY = {
    _CATEGORY_DETAIL: "细化",
    _CATEGORY_DESIGN: "设计",
}
_PRICING_SHEETS = {
    "设计定价": _CATEGORY_DESIGN,
    "细化定价": _CATEGORY_DETAIL,
}
_SETTLEMENT_HEADERS = [
    "行号",
    "名称",
    "项目号",
    "耗用数量",
    "单位",
    "单价(美元)",
    "总价(美元)",
    "折减系数",
    "结算代码",
    "类别",
    "结算年月",
    "发图年月",
]
_DEFAULT_DATA_DIR = (
    Path(__file__).resolve().parents[3]
    / "itss-steeltech-projects"
    / "生产计划"
    / "发图计划及预算结算"
)


def _now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _normalize_category(value: Any) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    return _CATEGORY_LABELS.get(text)


def _build_record_id(year: int, month: int, category: str, seq: int) -> str:
    return f"{year}_{month:02d}_{category}_{seq:03d}"


def _build_pending_id(category: str, seq: int) -> str:
    return f"pending_{category}_{seq:03d}"


def _normalize_project_no(value: Any) -> str:
    if value is None or value == "":
        return ""
    text = str(value).strip().upper()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def _build_pricing_id(category: str, pricing_code: str) -> str:
    return f"{category}_{pricing_code}"


def _normalize_float(value: Any, default: float = 0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _normalize_year_month(value: Any) -> tuple[int | None, int | None]:
    if value is None or value == "":
        return None, None
    if isinstance(value, datetime):
        return value.year, value.month
    if isinstance(value, date):
        return value.year, value.month
    text = str(value).strip()
    if not text:
        return None, None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m", "%Y/%m", "%Y.%m"):
        try:
            parsed = datetime.strptime(text, fmt)
            return parsed.year, parsed.month
        except ValueError:
            continue
    match = re.match(r"(\d{4})[-/.年](\d{1,2})", text)
    if match:
        return int(match.group(1)), int(match.group(2))
    return None, None


def _calc_total_price(quantity: float, unit_price: float, coefficient: float) -> float:
    return round(quantity * unit_price * coefficient, 4)


def list_pricing_basis(*, category: str | None = None) -> list[dict]:
    query = SettlementPricingBasis.query
    if category:
        query = query.filter_by(category=category)
    rows = query.order_by(
        SettlementPricingBasis.category.asc(),
        SettlementPricingBasis.seq.asc(),
        SettlementPricingBasis.pricing_code.asc(),
    ).all()
    return [row.to_dict() for row in rows]


def list_settlement_records(
    *,
    year: int | None = None,
    month: int | None = None,
    category: str | None = None,
) -> list[dict]:
    query = SettlementRecord.query.filter(
        SettlementRecord.settlement_year.isnot(None),
        SettlementRecord.settlement_month.isnot(None),
    )
    if year is not None:
        query = query.filter_by(settlement_year=year)
    if month is not None:
        query = query.filter_by(settlement_month=month)
    if category:
        query = query.filter_by(category=category)
    rows = query.order_by(
        SettlementRecord.settlement_year.desc(),
        SettlementRecord.settlement_month.desc(),
        SettlementRecord.category.asc(),
        SettlementRecord.seq.asc(),
    ).all()
    return [row.to_dict() for row in rows]


def list_pending_items(*, category: str | None = None) -> list[dict]:
    query = SettlementPendingItem.query
    if category:
        query = query.filter_by(category=category)
    rows = query.order_by(
        SettlementPendingItem.issue_year.desc(),
        SettlementPendingItem.issue_month.desc(),
        SettlementPendingItem.category.asc(),
        SettlementPendingItem.seq.asc(),
    ).all()
    return [row.to_dict() for row in rows]


def attach_pending_items_to_month(
    *,
    year: int,
    month: int,
    category: str,
    pending_ids: list[str] | None = None,
) -> tuple[list[dict], str | None, int]:
    """把未结算项目加入当月结算（成功后会把对应未结算行删除）。"""
    category_norm = _normalize_category(category)
    if not category_norm:
        return [], "category 无效", 400
    if month < 1 or month > 12:
        return [], "月份必须在 1-12 之间", 400

    query = SettlementPendingItem.query.filter_by(category=category_norm)
    if pending_ids:
        query = query.filter(SettlementPendingItem.id.in_(pending_ids))
    pending_rows = query.order_by(SettlementPendingItem.seq.asc()).all()
    if not pending_rows:
        return [], "未找到对应未结算项目", 400

    existing_count = (
        SettlementRecord.query.filter_by(
            settlement_year=year,
            settlement_month=month,
            category=category_norm,
        ).count()
    )

    now = _now_str()
    inserted: list[SettlementRecord] = []
    moved_ids = [row.id for row in pending_rows]

    for index, item in enumerate(pending_rows, start=1):
        seq = existing_count + index
        row = SettlementRecord(
            id=_build_record_id(year, month, category_norm, seq),
            year=year,
            month=month,
            seq=seq,
            category=category_norm,
            name=item.name,
            project_no=item.project_no,
            quantity=item.quantity,
            unit=item.unit,
            unit_price=item.unit_price,
            total_price=item.total_price,
            coefficient=item.coefficient,
            pricing_code=item.pricing_code,
            settlement_year=year,
            settlement_month=month,
            issue_year=item.issue_year,
            issue_month=item.issue_month,
            created_at=now,
            updated_at=now,
        )
        db.session.add(row)
        inserted.append(row)

    if pending_ids:
        SettlementPendingItem.query.filter(SettlementPendingItem.id.in_(moved_ids)).delete(
            synchronize_session=False
        )
    else:
        SettlementPendingItem.query.filter_by(category=category_norm).delete(
            synchronize_session=False
        )

    db.session.commit()
    return [row.to_dict() for row in inserted], None, 200


def move_month_records_to_pending(
    *,
    year: int,
    month: int,
    category: str,
    record_ids: list[str] | None = None,
) -> tuple[list[dict], str | None, int]:
    """把当月结算记录移回未结算项目（成功后会把对应结算行删除）。"""
    category_norm = _normalize_category(category)
    if not category_norm:
        return [], "category 无效", 400
    if month < 1 or month > 12:
        return [], "月份必须在 1-12 之间", 400

    query = SettlementRecord.query.filter_by(
        settlement_year=year, settlement_month=month, category=category_norm
    )
    if record_ids:
        query = query.filter(SettlementRecord.id.in_(record_ids))
    month_rows = query.order_by(SettlementRecord.seq.asc()).all()
    if not month_rows:
        return [], "未找到对应当月结算记录", 400

    existing_count = (
        SettlementPendingItem.query.filter_by(category=category_norm).count()
    )
    now = _now_str()
    inserted: list[SettlementPendingItem] = []
    moved_ids = [row.id for row in month_rows]

    for index, item in enumerate(month_rows, start=1):
        seq = existing_count + index
        row = SettlementPendingItem(
            id=_build_pending_id(category_norm, seq),
            category=category_norm,
            seq=seq,
            name=item.name,
            project_no=item.project_no,
            quantity=item.quantity,
            unit=item.unit,
            unit_price=item.unit_price,
            total_price=item.total_price,
            coefficient=item.coefficient,
            pricing_code=item.pricing_code,
            issue_year=item.issue_year,
            issue_month=item.issue_month,
            created_at=now,
            updated_at=now,
        )
        db.session.add(row)
        inserted.append(row)

    if record_ids:
        SettlementRecord.query.filter(SettlementRecord.id.in_(moved_ids)).delete(
            synchronize_session=False
        )
    else:
        SettlementRecord.query.filter_by(
            settlement_year=year, settlement_month=month, category=category_norm
        ).delete(synchronize_session=False)

    db.session.commit()
    return [row.to_dict() for row in inserted], None, 200


def get_pricing_by_code(*, category: str, pricing_code: str) -> dict | None:
    code = (pricing_code or "").strip().upper()
    if not code:
        return None
    row = SettlementPricingBasis.query.filter_by(category=category, pricing_code=code).first()
    return row.to_dict() if row else None


def batch_replace_month_records(records: list[dict]) -> tuple[list[dict], str | None, int]:
    if not records:
        return [], "请求体不能为空", 400

    first = records[0]
    year = first.get("year")
    month = first.get("month")
    category = _normalize_category(first.get("category"))
    if year is None or month is None or not category:
        return [], "每条记录必须包含 year、month、category", 400

    try:
        year = int(year)
        month = int(month)
    except (TypeError, ValueError):
        return [], "year 和 month 必须为整数", 400

    if month < 1 or month > 12:
        return [], "月份必须在 1-12 之间", 400

    for item in records:
        item_category = _normalize_category(item.get("category"))
        if item.get("year") != year or item.get("month") != month or item_category != category:
            return [], "批量保存的记录必须属于同一年月与类别", 400
        if not str(item.get("name") or "").strip():
            return [], "名称不能为空", 400

    SettlementRecord.query.filter_by(
        settlement_year=year,
        settlement_month=month,
        category=category,
    ).delete()

    now = _now_str()
    saved: list[SettlementRecord] = []
    for index, item in enumerate(records, start=1):
        seq = int(item.get("seq") or index)
        quantity = _normalize_float(item.get("quantity"))
        unit_price = _normalize_float(item.get("unitPrice"))
        coefficient = _normalize_float(item.get("coefficient"), default=1) or 1
        total_price = _normalize_float(item.get("totalPrice"))
        if total_price == 0 and quantity and unit_price:
            total_price = _calc_total_price(quantity, unit_price, coefficient)

        settlement_year = item.get("settlementYear")
        settlement_month = item.get("settlementMonth")
        issue_year = item.get("issueYear")
        issue_month = item.get("issueMonth")
        if settlement_year is None or settlement_month is None:
            settlement_year, settlement_month = year, month
        if issue_year is None or issue_month is None:
            issue_year, issue_month = settlement_year, settlement_month

        row = SettlementRecord(
            id=_build_record_id(year, month, category, seq),
            year=year,
            month=month,
            seq=seq,
            category=category,
            name=str(item.get("name") or "").strip(),
            project_no=str(item.get("projectNo") or "").strip().upper() or None,
            quantity=quantity,
            unit=str(item.get("unit") or "").strip(),
            unit_price=unit_price,
            total_price=total_price,
            coefficient=coefficient,
            pricing_code=str(item.get("pricingCode") or "").strip().upper(),
            settlement_year=int(settlement_year) if settlement_year is not None else None,
            settlement_month=int(settlement_month) if settlement_month is not None else None,
            issue_year=int(issue_year) if issue_year is not None else None,
            issue_month=int(issue_month) if issue_month is not None else None,
            created_at=now,
            updated_at=now,
        )
        db.session.add(row)
        saved.append(row)

    db.session.commit()
    return [row.to_dict() for row in saved], None, 200


def batch_replace_pending_items(records: list[dict]) -> tuple[list[dict], str | None, int]:
    if not records:
        return [], "请求体不能为空", 400

    category = _normalize_category(records[0].get("category"))
    if not category:
        return [], "每条记录必须包含 category", 400

    for item in records:
        item_category = _normalize_category(item.get("category"))
        if item_category != category:
            return [], "批量保存的记录必须属于同一类别", 400
        if not str(item.get("name") or "").strip():
            return [], "名称不能为空", 400

    SettlementPendingItem.query.filter_by(category=category).delete()

    now = _now_str()
    saved: list[SettlementPendingItem] = []
    for index, item in enumerate(records, start=1):
        seq = int(item.get("seq") or index)
        quantity = _normalize_float(item.get("quantity"))
        unit_price = _normalize_float(item.get("unitPrice"))
        coefficient = _normalize_float(item.get("coefficient"), default=1) or 1
        total_price = _normalize_float(item.get("totalPrice"))
        if total_price == 0 and quantity and unit_price:
            total_price = _calc_total_price(quantity, unit_price, coefficient)

        row = SettlementPendingItem(
            id=str(item.get("id") or _build_pending_id(category, seq)),
            category=category,
            seq=seq,
            name=str(item.get("name") or "").strip(),
            project_no=_normalize_project_no(item.get("projectNo")) or None,
            quantity=quantity,
            unit=str(item.get("unit") or "").strip(),
            unit_price=unit_price,
            total_price=total_price,
            coefficient=coefficient,
            pricing_code=str(item.get("pricingCode") or "").strip().upper(),
            issue_year=int(item["issueYear"]) if item.get("issueYear") is not None else None,
            issue_month=int(item["issueMonth"]) if item.get("issueMonth") is not None else None,
            created_at=now,
            updated_at=now,
        )
        db.session.add(row)
        saved.append(row)

    db.session.commit()
    return [row.to_dict() for row in saved], None, 200


def import_pricing_basis_excel(file_bytes: bytes) -> tuple[list[dict], str | None, int]:
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        return [], f"无法读取 Excel: {exc}", 400

    now = _now_str()
    imported: list[SettlementPricingBasis] = []

    for sheet_name, category in _PRICING_SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        worksheet = workbook[sheet_name]
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(cell or "").strip() for cell in rows[0]]
        code_idx = next((i for i, h in enumerate(header) if "代码" in h), 0)
        name_idx = next((i for i, h in enumerate(header) if h == "项目"), 1)
        calc_idx = next((i for i, h in enumerate(header) if "计算" in h), -1)
        price_idx = next((i for i, h in enumerate(header) if "单价" in h), 2 if category == _CATEGORY_DETAIL else 3)
        unit_idx = next((i for i, h in enumerate(header) if h == "单位"), price_idx + 1)
        remark_idx = next((i for i, h in enumerate(header) if h == "备注"), unit_idx + 1)

        SettlementPricingBasis.query.filter_by(category=category).delete()

        seq = 0
        for row in rows[1:]:
            if not row:
                continue
            pricing_code = str(row[code_idx] or "").strip().upper()
            item_name = str(row[name_idx] or "").strip()
            if not pricing_code or not item_name:
                continue
            seq += 1
            calc_process = str(row[calc_idx] or "").strip() if calc_idx >= 0 else ""
            unit_price = _normalize_float(row[price_idx])
            unit = str(row[unit_idx] or "").strip()
            remark = str(row[remark_idx] or "").strip() if remark_idx < len(row) else ""

            item = SettlementPricingBasis(
                id=_build_pricing_id(category, pricing_code),
                category=category,
                pricing_code=pricing_code,
                item_name=item_name,
                calc_process=calc_process or None,
                unit_price=unit_price,
                unit=unit,
                remark=remark or None,
                seq=seq,
                created_at=now,
                updated_at=now,
            )
            db.session.add(item)
            imported.append(item)

    db.session.commit()
    return [row.to_dict() for row in imported], None, 200


def import_settlement_records_excel(
    file_bytes: bytes,
    *,
    year: int | None = None,
    month: int | None = None,
    category: str | None = None,
) -> tuple[list[dict], str | None, int]:
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as exc:
        return [], f"无法读取 Excel: {exc}", 400

    sheet_name = next((name for name in workbook.sheetnames if "结算" in name), workbook.sheetnames[0])
    worksheet = workbook[sheet_name]
    rows = list(worksheet.iter_rows(values_only=True))
    if len(rows) < 2:
        return [], "Excel 中没有有效数据", 400

    header = [str(cell or "").strip() for cell in rows[0]]
    col = {name: idx for idx, name in enumerate(header) if name}

    def get_col(*candidates: str, default: int = -1) -> int:
        for candidate in candidates:
            if candidate in col:
                return col[candidate]
        return default

    idx_seq = get_col("行号", default=0)
    idx_name = get_col("名称", default=1)
    idx_project_no = get_col("项目号", default=2)
    idx_quantity = get_col("耗用数量", "结算数量", default=3)
    idx_unit = get_col("单位", default=4)
    idx_unit_price = get_col("单价(美元)", "单价(万元)", default=5)
    idx_total_price = get_col("总价(美元)", "总价(万元)", default=6)
    idx_coefficient = get_col("折减系数", "累计系数", default=7)
    idx_pricing_code = get_col("结算代码", "定价代码", default=8)
    idx_category = get_col("类别", default=9)
    idx_settlement_ym = get_col("结算年月", "结算日期", default=10)
    idx_issue_ym = get_col("发图年月", "发图日期", default=11)

    grouped: dict[tuple[int, int, str], list[dict]] = {}
    pending_by_category: dict[str, list[dict]] = {}
    for row in rows[1:]:
        if not row or not any(row):
            continue
        name = str(row[idx_name] or "").strip() if idx_name >= 0 else ""
        if not name:
            continue

        row_category = _normalize_category(row[idx_category] if idx_category >= 0 else category)
        if not row_category:
            continue
        if category and row_category != _normalize_category(category):
            continue

        settlement_year, settlement_month = _normalize_year_month(row[idx_settlement_ym] if idx_settlement_ym >= 0 else None)
        issue_year, issue_month = _normalize_year_month(row[idx_issue_ym] if idx_issue_ym >= 0 else None)

        quantity = _normalize_float(row[idx_quantity] if idx_quantity >= 0 else 0)
        unit_price = _normalize_float(row[idx_unit_price] if idx_unit_price >= 0 else 0)
        coefficient = _normalize_float(row[idx_coefficient] if idx_coefficient >= 0 else 1, default=1) or 1
        total_price = _normalize_float(row[idx_total_price] if idx_total_price >= 0 else 0)
        if total_price == 0:
            total_price = _calc_total_price(quantity, unit_price, coefficient)

        common = {
            "name": name,
            "projectNo": _normalize_project_no(row[idx_project_no] if idx_project_no >= 0 else ""),
            "quantity": quantity,
            "unit": str(row[idx_unit] or "").strip() if idx_unit >= 0 else "",
            "unitPrice": unit_price,
            "totalPrice": total_price,
            "coefficient": coefficient,
            "pricingCode": str(row[idx_pricing_code] or "").strip().upper() if idx_pricing_code >= 0 else "",
            "issueYear": issue_year,
            "issueMonth": issue_month,
        }

        if settlement_year is None or settlement_month is None:
            pending_by_category.setdefault(row_category, []).append(
                {
                    **common,
                    "category": row_category,
                    "seq": int(_normalize_float(row[idx_seq] if idx_seq >= 0 else len(pending_by_category[row_category]) + 1)),
                }
            )
            continue

        record_year = year or settlement_year
        record_month = month or settlement_month
        key = (record_year, record_month, row_category)
        grouped.setdefault(key, []).append(
            {
                "year": record_year,
                "month": record_month,
                "seq": int(_normalize_float(row[idx_seq] if idx_seq >= 0 else len(grouped[key]) + 1)),
                "category": row_category,
                **common,
                "settlementYear": settlement_year,
                "settlementMonth": settlement_month,
            }
        )

    if not grouped and not pending_by_category:
        return [], "未解析到有效结算记录", 400

    all_saved: list[dict] = []
    for _, items in sorted(grouped.items()):
        items.sort(key=lambda item: item["seq"])
        for index, item in enumerate(items, start=1):
            item["seq"] = index
        saved, error, status = batch_replace_month_records(items)
        if error:
            return [], error, status
        all_saved.extend(saved)

    for pending_category, items in pending_by_category.items():
        items.sort(key=lambda item: item["seq"])
        for index, item in enumerate(items, start=1):
            item["seq"] = index
        saved, error, status = batch_replace_pending_items(items)
        if error:
            return [], error, status
        all_saved.extend(saved)

    return all_saved, None, 200


def export_settlement_records_excel(year: int, month: int, category: str) -> bytes:
    category = _normalize_category(category) or _CATEGORY_DETAIL
    records = list_settlement_records(year=year, month=month, category=category)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "结算表"

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center")

    for col_index, header in enumerate(_SETTLEMENT_HEADERS, start=1):
        cell = worksheet.cell(row=1, column=col_index, value=header)
        cell.font = header_font
        cell.alignment = center
        cell.border = border

    for row_index, record in enumerate(records, start=2):
        settlement_ym = None
        if record.get("settlementYear") and record.get("settlementMonth"):
            settlement_ym = datetime(record["settlementYear"], record["settlementMonth"], 1)
        issue_ym = None
        if record.get("issueYear") and record.get("issueMonth"):
            issue_ym = datetime(record["issueYear"], record["issueMonth"], 1)

        values = [
            record.get("seq"),
            record.get("name"),
            record.get("projectNo") or None,
            record.get("quantity"),
            record.get("unit"),
            record.get("unitPrice"),
            record.get("totalPrice"),
            record.get("coefficient"),
            record.get("pricingCode"),
            _CATEGORY_DISPLAY.get(record.get("category", ""), record.get("category")),
            settlement_ym,
            issue_ym,
        ]
        for col_index, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_index, column=col_index, value=value)
            cell.border = border
            if col_index in (11, 12) and isinstance(value, datetime):
                cell.number_format = "yyyy-mm"

    for col_index in range(1, len(_SETTLEMENT_HEADERS) + 1):
        worksheet.column_dimensions[chr(64 + col_index)].width = 14
    worksheet.column_dimensions["B"].width = 42

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def seed_settlement_data_if_empty() -> None:
    pricing_path = _DEFAULT_DATA_DIR / "钢结构计划科定价表.xlsx"
    settlement_path = _DEFAULT_DATA_DIR / "结算记录.xlsx"

    if SettlementPricingBasis.query.count() == 0 and pricing_path.is_file():
        with pricing_path.open("rb") as handle:
            import_pricing_basis_excel(handle.read())

    should_import_settlement = (
        settlement_path.is_file()
        and (
            SettlementRecord.query.count() == 0
            or SettlementPendingItem.query.count() == 0
        )
    )
    if should_import_settlement:
        with settlement_path.open("rb") as handle:
            import_settlement_records_excel(handle.read())
